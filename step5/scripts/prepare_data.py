#!/usr/bin/env python3
"""Export Step 5 test inputs from the committed golden-set workbook.

Writes (all derived only from team 9 rows 9000-9999):
  step5/data/tickets_load.tsv  row<TAB>JSON request body, for JMeter load traffic.
                               Uses the 840 non-golden rows (9000-9839) so load
                               traffic and accuracy traffic stay separate.
  step5/data/golden_set.csv    row,label,narrative for the 160 frozen golden tickets.

Usage: python step5/scripts/prepare_data.py
Requires: openpyxl
"""
import csv
import json
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
WORKBOOK = ROOT / "ICT3113 Golden Set P1-9.xlsx"
OUT = ROOT / "step5" / "data"
CATEGORIES = {
    "Credit reporting",
    "Debt collection",
    "Mortgage",
    "Credit card",
    "Bank account or service",
    "Consumer loan",
    "Money transfer or service",
}


def main() -> int:
    wb = openpyxl.load_workbook(WORKBOOK, read_only=True, data_only=True)
    team_rows = {
        int(row): narrative
        for row, _label, narrative in wb["P1-9 Data"].iter_rows(min_row=2, max_col=3, values_only=True)
        if row is not None
    }
    golden = [
        (int(row), label, narrative)
        for row, label, narrative, *_ in wb["Just Data"].iter_rows(min_row=2, max_col=3, values_only=True)
        if row is not None
    ]

    assert len(team_rows) == 1000 and min(team_rows) == 9000 and max(team_rows) == 9999
    assert len(golden) == 160 and len({r for r, _, _ in golden}) == 160
    assert all(label in CATEGORIES for _, label, _ in golden)
    assert all(team_rows[r] == n for r, _, n in golden), "golden narrative differs from team rows"

    golden_rows = {r for r, _, _ in golden}
    OUT.mkdir(parents=True, exist_ok=True)

    with (OUT / "tickets_load.tsv").open("w", encoding="utf-8", newline="\n") as f:
        f.write("row\tbody\n")
        count = 0
        for row in sorted(team_rows):
            if row in golden_rows:
                continue
            # json.dumps escapes quotes, tabs and newlines, so each body is one TSV field.
            f.write(f"{row}\t{json.dumps({'narrative': team_rows[row]}, ensure_ascii=True)}\n")
            count += 1

    with (OUT / "golden_set.csv").open("w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["row", "label", "narrative"])
        writer.writerows(golden)

    print(f"tickets_load.tsv: {count} rows; golden_set.csv: {len(golden)} rows")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
